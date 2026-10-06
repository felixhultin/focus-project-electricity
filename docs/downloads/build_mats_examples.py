#!/usr/bin/env python3
"""Build a browsable, paginated list of examples from data/mats.

Requires openpyxl. The output is static JSON and HTML for GitHub Pages.
Rows without an HTTP(S) source URL are included as plain text.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import Counter, OrderedDict
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
HEADER = ("token", "lemma", "pos", "id", "date", "url", "target", "start", "end", "text")
CHUNK_SIZE = 500
MAX_OPEN_FILES = 48
YEAR = re.compile(r"^(\d{4})(?:-|$)")
SPACE = re.compile(r"\s+")


def is_source_url(value: object) -> bool:
    if not isinstance(value, str):
        return False
    parsed = urlsplit(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def excerpt(text: object, token: object, start: object, end: object) -> tuple[str, str, str]:
    content = str(text or "")
    word = str(token or "")
    try:
        left, right = int(start), int(end)
    except (ValueError, TypeError):
        left = right = -1
    if not (0 <= left < right <= len(content)):
        left = content.casefold().find(word.casefold()) if word else -1
        right = left + len(word) if left >= 0 else -1
    if left < 0:
        return "", "", SPACE.sub(" ", content[:220]).strip()
    before = SPACE.sub(" ", content[max(0, left - 110):left])
    match = SPACE.sub(" ", content[left:right])
    after = SPACE.sub(" ", content[right:right + 110])
    return (
        ("…" if left > 110 else "") + before,
        match,
        after + ("…" if right + 110 < len(content) else ""),
    )


class ChunkWriter:
    def __init__(self, root: Path):
        self.root = root
        self.counts = Counter()
        self.handles: OrderedDict[Path, object] = OrderedDict()

    def add(self, key: tuple[str, str, str], row: list[object]) -> None:
        count = self.counts[key]
        page = count // CHUNK_SIZE
        path = self.root.joinpath(*key, f"{page}.json")
        if path in self.handles:
            stream = self.handles.pop(path)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            stream = path.open("w" if count % CHUNK_SIZE == 0 else "a", encoding="utf-8")
            if count % CHUNK_SIZE == 0:
                stream.write("[")
        stream.write(("," if count % CHUNK_SIZE else "") + json.dumps(row, ensure_ascii=False, separators=(",", ":")))
        self.counts[key] += 1
        if self.counts[key] % CHUNK_SIZE == 0:
            stream.write("]")
            stream.close()
        else:
            self.handles[path] = stream
            if len(self.handles) > MAX_OPEN_FILES:
                _, old = self.handles.popitem(last=False)
                old.close()

    def close(self) -> None:
        for stream in self.handles.values():
            stream.close()
        self.handles.clear()
        for key, count in self.counts.items():
            if count % CHUNK_SIZE:
                path = self.root.joinpath(*key, f"{count // CHUNK_SIZE}.json")
                with path.open("a", encoding="utf-8") as stream:
                    stream.write("]")


def build(source: Path, site: Path) -> dict[str, object]:
    destination = site / "example_data"
    if destination.exists():
        shutil.rmtree(destination)
    destination.mkdir(parents=True)
    writer = ChunkWriter(destination)
    corpora: list[str] = []
    corpus_ids: dict[str, int] = {}
    all_rows = Counter()
    missing = Counter()
    linked_counts = Counter()
    words = set()
    try:
        for resource in ("lb", "tkr", "kubhist2"):
            paths = sorted((source / resource).glob("*.xlsx"))
            if not paths:
                raise ValueError(f"No workbooks in {source / resource}")
            for path in paths:
                corpus, word = path.stem.rsplit("_", 1)
                words.add(word)
                if corpus not in corpus_ids:
                    corpus_ids[corpus] = len(corpora)
                    corpora.append(corpus)
                corpus_id = corpus_ids[corpus]
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    sheet = workbook.active
                    rows = sheet.iter_rows(values_only=True)
                    if tuple(next(rows, ())) != HEADER:
                        raise ValueError(f"Unexpected columns in {path}")
                    for row in rows:
                        all_rows[resource] += 1
                        url = row[5] if is_source_url(row[5]) else None
                        if url is None:
                            missing[resource] += 1
                        date = str(row[4] or "")
                        found_year = YEAR.match(date)
                        decade = str(int(found_year.group(1)) // 10 * 10) if found_year else "undated"
                        before, hit, after = excerpt(row[9], row[0], row[7], row[8])
                        key = (word, resource, decade)
                        writer.add(key, [date, corpus_id, before, hit, after, url])
                        if url is not None:
                            linked_counts[key] += 1
                finally:
                    workbook.close()
    finally:
        writer.close()

    manifest = {
        "chunk_size": CHUNK_SIZE,
        "words": sorted(words),
        "resources": ["lb", "tkr", "kubhist2"],
        "corpora": corpora,
        "all_rows": dict(all_rows),
        "without_source_url": dict(missing),
        "buckets": [
            {"word": word, "resource": resource, "decade": decade, "count": count,
             "with_source_url": linked_counts[(word, resource, decade)]}
            for (word, resource, decade), count in sorted(writer.counts.items())
        ],
    }
    (destination / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    shutil.copy2(ROOT / "mats_examples_viewer.html", site / "examples.html")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "data" / "mats")
    parser.add_argument("--site", type=Path, default=ROOT / "korp-frequency-pages" / "docs")
    args = parser.parse_args()
    args.site.mkdir(parents=True, exist_ok=True)
    manifest = build(args.source, args.site)
    linked = sum(bucket["with_source_url"] for bucket in manifest["buckets"])
    total = sum(manifest["all_rows"].values())
    print(f"Built {args.site / 'examples.html'} with {total:,} examples; {linked:,} have source URLs")


if __name__ == "__main__":
    main()
